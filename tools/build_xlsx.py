"""Usage: python3 build_xlsx.py posts.json batch-NN image|text [YYYY-MM-DD]"""
import json, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
posts, batch, kind = json.load(open(sys.argv[1])), sys.argv[2], sys.argv[3]
day = sys.argv[4] if len(sys.argv) > 4 else datetime.date.today().isoformat()
base = f'https://raw.githubusercontent.com/anonymouscoder3-jpg/just-money-assets/main/posts/{batch}/'
wb = Workbook(); ws = wb.active; ws.title = 'Sheet1'
ws['A1'] = '## DO NOT EDIT OR DELETE THIS ROW'; ws['B1'] = 'BULK UPLOAD VERSION 3.1'
ws['A2'] = '## Rows starting with ## are ignored. One post per row from row 5. Leave POSTING TIME empty to add posts to the Posting Plan.'
hdr = ['POSTING TIME','CAPTION','LINK','MEDIA','TITLE','FIRST COMMENT','FIRST COMMENT DELAY','TOBI BACKGROUND','REEL OR STORY']
for i, h in enumerate(hdr):
    c = ws.cell(3, i + 1, ('## ' if i == 0 else '') + h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F2937')
notes = ['## MM/DD/YYYY HH:MM (24h) or empty','Post text','URL (optional)','Public .jpg/.png URL','Optional','FB/IG/LinkedIn only','5, 10, 15, 30 or 60','FB text posts only','reel or story']
for i, n in enumerate(notes): ws.cell(4, i + 1, n).font = Font(italic=True, color='6B7280')
for r, p in enumerate(posts, start=5):
    ws.cell(r, 2, p['caption'] if kind == 'image' else p['text']).alignment = Alignment(wrap_text=True, vertical='top')
    if kind == 'image': ws.cell(r, 4, base + p['file'])
    if p.get('time'): ws.cell(r, 1, p['time'])
for col, w in zip('ABCDEFGHI', [22, 70, 14, 95, 12, 18, 12, 14, 12]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'A5'
out = f'justmoney_{kind}_posts_{day}.xlsx'; wb.save(out); print(out)
